"""Holdings template importer and snapshot normalization."""

from __future__ import annotations

import csv
from dataclasses import asdict, dataclass
from pathlib import Path

ALLOWED_ACCOUNT_TYPES = {"taxable", "irp", "pension"}
REQUIRED_COLUMNS = {
    "account_id",
    "account_type",
    "symbol",
    "name",
    "qty",
    "currency",
    "price",
    "value",
}


class HoldingsImportError(ValueError):
    """Raised when holdings import data is invalid."""


@dataclass(frozen=True)
class PositionSnapshot:
    """Normalized position row."""

    account_id: str
    account_type: str
    symbol: str
    name: str
    qty: float
    currency: str
    price: float | None
    value: float | None


@dataclass(frozen=True)
class CashSnapshot:
    """Normalized cash row."""

    account_id: str
    account_type: str
    currency: str
    amount: float


def _coerce_float(value: object) -> float | None:
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)

    text = str(value).strip()
    if text == "":
        return None

    normalized = text.replace(",", "")
    try:
        return float(normalized)
    except ValueError as exc:
        raise HoldingsImportError(f"Could not parse numeric value: {value!r}") from exc


def _normalize_text(value: object) -> str:
    return str(value or "").strip()


def _is_cash_row(symbol: str, name: str) -> bool:
    symbol_upper = symbol.upper()
    name_lower = name.lower()
    return symbol_upper in {"CASH", "현금"} or name_lower in {"cash", "현금"}


def _rows_from_csv(path: Path) -> list[dict[str, object]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        if reader.fieldnames is None:
            raise HoldingsImportError("CSV file has no header row.")
        return [dict(row) for row in reader]


def _rows_from_xlsx(path: Path) -> list[dict[str, object]]:
    try:
        from openpyxl import load_workbook
    except ModuleNotFoundError as exc:
        raise HoldingsImportError(
            "openpyxl is required to import .xlsx files. Install dependencies first."
        ) from exc

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        values = sheet.iter_rows(values_only=True)
        header = next(values, None)
        if header is None:
            raise HoldingsImportError("XLSX file has no header row.")
        columns = [str(item or "").strip() for item in header]
        rows: list[dict[str, object]] = []
        for row_values in values:
            row_dict = {
                columns[idx]: row_values[idx] if idx < len(row_values) else None
                for idx in range(len(columns))
            }
            rows.append(row_dict)
        return rows
    finally:
        workbook.close()


def read_holdings_rows(path: Path) -> list[dict[str, object]]:
    """Read holdings template rows from CSV or XLSX."""
    suffix = path.suffix.lower()
    if suffix == ".csv":
        rows = _rows_from_csv(path)
    elif suffix == ".xlsx":
        rows = _rows_from_xlsx(path)
    else:
        raise HoldingsImportError("Unsupported input format. Use .csv or .xlsx")

    if not rows:
        raise HoldingsImportError("Input file contains no data rows.")

    row_columns = {key for key in rows[0] if key}
    missing = sorted(REQUIRED_COLUMNS - row_columns)
    if missing:
        raise HoldingsImportError(f"Missing required columns: {', '.join(missing)}")

    return rows


def normalize_snapshots(
    rows: list[dict[str, object]],
) -> tuple[list[PositionSnapshot], list[CashSnapshot]]:
    """Normalize holdings rows into position and cash snapshots."""
    positions: list[PositionSnapshot] = []
    cash: list[CashSnapshot] = []

    for index, row in enumerate(rows, start=2):
        account_id = _normalize_text(row.get("account_id"))
        account_type = _normalize_text(row.get("account_type")).lower()
        symbol = _normalize_text(row.get("symbol"))
        name = _normalize_text(row.get("name"))
        currency = _normalize_text(row.get("currency")).upper()
        qty = _coerce_float(row.get("qty"))
        price = _coerce_float(row.get("price"))
        value = _coerce_float(row.get("value"))

        if not account_id:
            raise HoldingsImportError(f"Row {index}: account_id is required")
        if account_type not in ALLOWED_ACCOUNT_TYPES:
            raise HoldingsImportError(
                f"Row {index}: account_type must be one of {sorted(ALLOWED_ACCOUNT_TYPES)}"
            )
        if not currency:
            raise HoldingsImportError(f"Row {index}: currency is required")
        if qty is None:
            raise HoldingsImportError(f"Row {index}: qty is required")

        computed_value = value
        if computed_value is None and price is not None:
            computed_value = qty * price

        if _is_cash_row(symbol=symbol, name=name):
            cash_amount = computed_value if computed_value is not None else qty
            cash.append(
                CashSnapshot(
                    account_id=account_id,
                    account_type=account_type,
                    currency=currency,
                    amount=cash_amount,
                )
            )
            continue

        if not symbol:
            raise HoldingsImportError(f"Row {index}: symbol is required for non-cash rows")
        if not name:
            raise HoldingsImportError(f"Row {index}: name is required for non-cash rows")

        positions.append(
            PositionSnapshot(
                account_id=account_id,
                account_type=account_type,
                symbol=symbol,
                name=name,
                qty=qty,
                currency=currency,
                price=price,
                value=computed_value,
            )
        )

    return positions, cash


def write_snapshots_parquet(
    positions: list[PositionSnapshot],
    cash: list[CashSnapshot],
    out_dir: Path,
) -> tuple[Path, Path]:
    """Write normalized snapshots as parquet files."""
    try:
        import pyarrow as pa
        import pyarrow.parquet as pq
    except ModuleNotFoundError as exc:
        raise HoldingsImportError(
            "pyarrow is required to write parquet outputs. Install dependencies first."
        ) from exc

    out_dir.mkdir(parents=True, exist_ok=True)
    positions_path = out_dir / "positions_snapshot.parquet"
    cash_path = out_dir / "cash_snapshot.parquet"

    positions_schema = pa.schema(
        [
            ("account_id", pa.string()),
            ("account_type", pa.string()),
            ("symbol", pa.string()),
            ("name", pa.string()),
            ("qty", pa.float64()),
            ("currency", pa.string()),
            ("price", pa.float64()),
            ("value", pa.float64()),
        ]
    )
    cash_schema = pa.schema(
        [
            ("account_id", pa.string()),
            ("account_type", pa.string()),
            ("currency", pa.string()),
            ("amount", pa.float64()),
        ]
    )

    position_rows = [asdict(item) for item in positions]
    cash_rows = [asdict(item) for item in cash]

    positions_table = pa.Table.from_pylist(position_rows, schema=positions_schema)
    cash_table = pa.Table.from_pylist(cash_rows, schema=cash_schema)

    pq.write_table(positions_table, positions_path)
    pq.write_table(cash_table, cash_path)

    return positions_path, cash_path


def import_holdings(input_path: Path, out_dir: Path) -> tuple[Path, Path]:
    """End-to-end holdings template import pipeline."""
    rows = read_holdings_rows(input_path)
    positions, cash = normalize_snapshots(rows)
    return write_snapshots_parquet(positions=positions, cash=cash, out_dir=out_dir)
